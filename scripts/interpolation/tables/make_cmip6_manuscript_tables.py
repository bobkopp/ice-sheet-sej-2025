#!/usr/bin/env python3
"""Create final CMIP6 manuscript/SI tables from the adopted production chain.

Production source: emulandice-2050 + component-resolved temporal coupling + tension p=0.03.

Outputs
-------
1. Wide manuscript table (Table-3 style):
   rows = year -> ice-sheet component
   columns = CMIP6 scenario groups, each with median / 17–83 / 5–95

2. Long supplementary table:
   sections = ice-sheet component -> year -> scenario
   columns = scenario / median / 17–83 / 5–95

2050 is REQUIRED and is populated from the final selected-year source CSV.

Created on: 2026-10-26
Author: Elizabeth H. Davis, Ph.D. (email: elizabeth.h.davis@rutgers.edu)
"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

import pandas as pd

from table_config import (
    CMIP6_CODES,
    CMIP6_LABELS,
    CMIP6_OUTPUT_DIR,
    CMIP6_SOURCE,
    COMPONENTS,
    METHOD_NOTE_CMIP6,
    TABLE_YEARS,
)
from table_common import display_number, formatted_entries, load_projection_table

COMPONENT_DISPLAY = {
    "EAIS": "East Antarctic Ice Sheet (EAIS)",
    "WAIS": "West Antarctic Ice Sheet (WAIS)",
    "GrIS": "Greenland Ice Sheet (GrIS)",
    "Total AIS": "Total Antarctic Ice Sheet (AIS)",
    "Total ice sheet": "Total ice sheets",
}

WIDE_TITLE = (
    "Table 3. Projected contributions to global mean sea-level rise (cm) from "
    "the East Antarctic Ice Sheet (EAIS), West Antarctic Ice Sheet (WAIS), "
    "Greenland Ice Sheet (GrIS), total Antarctic Ice Sheet (AIS), and total ice "
    "sheets under CMIP6 SSP scenarios in 2050, 2100, 2150, 2200, and 2300."
)

LONG_TITLE = "CMIP6 ice-sheet contributions to global mean sea-level rise (cm)"


def build_wide_table(data: pd.DataFrame, decimals: int) -> pd.DataFrame:
    """CSV-friendly version of the wide manuscript table."""
    lookup = data.set_index(["ssp", "year", "component"])
    columns = ["Year and ice sheet"]
    for code in CMIP6_CODES:
        label = CMIP6_LABELS[code]
        columns.extend(
            [
                f"{label}: 50%",
                f"{label}: 17–83%",
                f"{label}: 5–95%",
            ]
        )

    rows = []
    for year in TABLE_YEARS:
        year_row = {column: "" for column in columns}
        year_row["Year and ice sheet"] = str(year)
        rows.append(year_row)
        for component in COMPONENTS:
            row = {"Year and ice sheet": component}
            for code in CMIP6_CODES:
                label = CMIP6_LABELS[code]
                q = lookup.loc[(code, year, component)]
                med, inner, outer = formatted_entries(q, decimals)
                row[f"{label}: 50%"] = med
                row[f"{label}: 17–83%"] = inner
                row[f"{label}: 5–95%"] = outer
            rows.append(row)
    return pd.DataFrame(rows, columns=columns)


def build_long_table(data: pd.DataFrame, decimals: int) -> pd.DataFrame:
    """Tidy CSV version of the long component/year/scenario table."""
    lookup = data.set_index(["component", "year", "ssp"])
    rows = []
    for component in COMPONENTS:
        for year in TABLE_YEARS:
            for code in CMIP6_CODES:
                med, inner, outer = formatted_entries(
                    lookup.loc[(component, year, code)], decimals
                )
                rows.append(
                    {
                        "ice_sheet": component,
                        "year": year,
                        "cmip6_scenario": CMIP6_LABELS[code],
                        "median_cm": med,
                        "p17_to_p83_cm": inner,
                        "p05_to_p95_cm": outer,
                    }
                )
    return pd.DataFrame(rows)


def _require_2050_populated(data: pd.DataFrame) -> None:
    """Fail loudly if 2050 is absent or has missing quantiles."""
    rows = data.loc[data["year"].eq(2050)].copy()
    expected = len(CMIP6_CODES) * len(COMPONENTS)
    if len(rows) != expected:
        raise ValueError(
            f"CMIP6: expected {expected} complete 2050 rows, found {len(rows)}. "
            "2050 will not be left blank."
        )
    qcols = ["p05_cm", "p17_cm", "p50_cm", "p83_cm", "p95_cm"]
    if rows[qcols].isna().any().any():
        raise ValueError("CMIP6: one or more 2050 quantiles are missing.")


def make_wide_xlsx(data: pd.DataFrame, path: Path, decimals: int) -> None:
    """Create Table-3-style wide workbook matching the supplied example."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "CMIP6 wide table"
    ws.sheet_view.showGridLines = False

    yellow = "FFF200"
    pale_yellow = "FFF4B3"
    gray = "E7E7E7"
    white = "FFFFFF"
    black = "000000"
    rule = "A6A6A6"
    thin = Side(style="thin", color=rule)
    hair = Side(style="hair", color=rule)

    ncols = 1 + 3 * len(CMIP6_CODES)
    last_col = get_column_letter(ncols)

    # Title row.
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title = ws.cell(1, 1, WIDE_TITLE)
    title.font = Font(name="Aptos", size=11, bold=False, color=black)
    title.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 45

    # Two-tier scenario header.
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)
    a = ws.cell(2, 1, "Year and ice\nsheet")
    a.font = Font(name="Aptos", size=9, bold=True)
    a.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    a.fill = PatternFill("solid", fgColor=gray)

    col = 2
    for code in CMIP6_CODES:
        label = CMIP6_LABELS[code]
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 2)
        cell = ws.cell(2, col, label)
        cell.font = Font(name="Aptos", size=9, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(col, col + 3):
            ws.cell(2, c).fill = PatternFill("solid", fgColor=gray)
        for offset, text in enumerate(("50%", "17–83%", "5–95%")):
            sub = ws.cell(3, col + offset, text)
            sub.font = Font(name="Aptos", size=8, bold=True)
            sub.alignment = Alignment(horizontal="center", vertical="center")
            sub.fill = PatternFill("solid", fgColor=gray)
        col += 3

    for c in range(1, ncols + 1):
        ws.cell(3, c).border = Border(bottom=Side(style="medium", color=black))

    lookup = data.set_index(["ssp", "year", "component"])
    row = 4
    for year in TABLE_YEARS:
        # Yellow year row, across full width, like the supplied example.
        ws.cell(row, 1, str(year))
        ws.cell(row, 1).font = Font(name="Aptos", size=9, bold=True)
        for c in range(1, ncols + 1):
            ws.cell(row, c).fill = PatternFill("solid", fgColor=yellow)
        ws.row_dimensions[row].height = 18
        row += 1

        for component in COMPONENTS:
            label = ws.cell(row, 1, component)
            label.font = Font(name="Aptos", size=9, italic=("Total" in component))
            label.fill = PatternFill("solid", fgColor=pale_yellow)
            label.alignment = Alignment(horizontal="left", vertical="center")

            col = 2
            for code in CMIP6_CODES:
                med, inner, outer = formatted_entries(
                    lookup.loc[(code, year, component)], decimals
                )
                for value in (med, inner, outer):
                    cell = ws.cell(row, col, value)
                    cell.font = Font(name="Aptos", size=8.5)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = Border(bottom=hair)
                    col += 1
            ws.row_dimensions[row].height = 17
            row += 1

    # Note.
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    note = ws.cell(row, 1, "Note. " + METHOD_NOTE_CMIP6)
    note.font = Font(name="Aptos", size=8, italic=True)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 60

    ws.column_dimensions["A"].width = 19
    for c in range(2, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 10.5

    ws.freeze_panes = "B4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:{last_col}{row}"
    wb.save(path)


def make_long_xlsx(data: pd.DataFrame, path: Path, decimals: int) -> None:
    """Create the long ice-sheet -> year -> scenario workbook matching screenshot 1."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "CMIP6 long table"
    ws.sheet_view.showGridLines = False

    gray = "F2F2F2"
    component_fill = "F7F7F7"
    rule = "A6A6A6"
    thin_bottom = Border(bottom=Side(style="thin", color=rule))
    medium_bottom = Border(bottom=Side(style="medium", color="000000"))

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 21
    ws.column_dimensions["D"].width = 21

    lookup = data.set_index(["component", "year", "ssp"])
    row = 1

    for comp_i, component in enumerate(COMPONENTS):
        if comp_i:
            row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row, 1, COMPONENT_DISPLAY[component])
        cell.font = Font(name="Aptos", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill("solid", fgColor=component_fill)
        ws.row_dimensions[row].height = 22
        row += 1

        headers = (
            "Year and CMIP6\nscenario",
            "Median\n(cm)",
            "17th–83rd\npercentile\n(cm)",
            "5th–95th\npercentile\n(cm)",
        )
        for c, text in enumerate(headers, start=1):
            h = ws.cell(row, c, text)
            h.font = Font(name="Aptos", size=9, bold=True)
            h.alignment = Alignment(horizontal="center" if c > 1 else "left", vertical="center", wrap_text=True)
            h.border = medium_bottom
        ws.row_dimensions[row].height = 40
        row += 1

        for year in TABLE_YEARS:
            # Year heading row.
            y = ws.cell(row, 1, str(year))
            y.font = Font(name="Aptos", size=9, bold=True)
            for c in range(1, 5):
                ws.cell(row, c).fill = PatternFill("solid", fgColor=gray)
            ws.row_dimensions[row].height = 18
            row += 1

            for code_i, code in enumerate(CMIP6_CODES):
                med, inner, outer = formatted_entries(
                    lookup.loc[(component, year, code)], decimals
                )
            
                values = (CMIP6_LABELS[code], med, inner, outer)
            
                for c, value in enumerate(values, start=1):
                    cell = ws.cell(row, c, value)
                    cell.font = Font(name="Aptos", size=9)
                    cell.alignment = Alignment(
                        horizontal="left" if c == 1 else "center",
                        vertical="center",
                        indent=2 if c == 1 else 0,
                    )
            
                    # Only draw a rule after the LAST SSP in each projection year.
                    if code_i == len(CMIP6_CODES) - 1:
                        cell.border = thin_bottom
            
                ws.row_dimensions[row].height = 18
                row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    note = ws.cell(row, 1, "Note. " + METHOD_NOTE_CMIP6)
    note.font = Font(name="Aptos", size=8, italic=True)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 65

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:D{row}"
    wb.save(path)


def run(*, source: Path = CMIP6_SOURCE, output_dir: Path = CMIP6_OUTPUT_DIR, decimals: int = 0):
    if decimals not in range(5):
        raise ValueError("decimals must be 0, 1, 2, 3, or 4")

    data = load_projection_table(
        source, scenario_codes=CMIP6_CODES, family="CMIP6"
    )
    _require_2050_populated(data)

    output_dir = Path(output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    # Full-precision audit.
    audit = output_dir / "CMIP6_quantiles_full_precision_cm.csv"
    data.to_csv(audit, index=False)

    # Wide manuscript layout — CSV + formatted XLSX.
    wide = build_wide_table(data, decimals)
    wide_csv = output_dir / "Table3_CMIP6_wide_manuscript_final.csv"
    wide_xlsx = output_dir / "Table3_CMIP6_wide_manuscript_final.xlsx"
    wide.to_csv(wide_csv, index=False, quoting=csv.QUOTE_MINIMAL)
    make_wide_xlsx(data, wide_xlsx, decimals)

    # Long SI layout — CSV + formatted XLSX.
    long = build_long_table(data, decimals)
    long_csv = output_dir / "SI_CMIP6_long_table_all_components_final.csv"
    long_xlsx = output_dir / "SI_CMIP6_long_table_all_components_final.xlsx"
    long.to_csv(long_csv, index=False)
    make_long_xlsx(data, long_xlsx, decimals)

    print("\nCMIP6 final-production tables complete")
    print("Source:", Path(source).expanduser().resolve())
    print("Wide CSV:", wide_csv)
    print("Wide XLSX:", wide_xlsx)
    print("Long CSV:", long_csv)
    print("Long XLSX:", long_xlsx)
    print("2050: explicitly validated and populated")
    print("Audit table:", audit)
    return {
        "wide_csv": wide_csv,
        "wide_xlsx": wide_xlsx,
        "long_csv": long_csv,
        "long_xlsx": long_xlsx,
        "audit": audit,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=CMIP6_SOURCE)
    parser.add_argument("--output-dir", type=Path, default=CMIP6_OUTPUT_DIR)
    parser.add_argument("--decimals", type=int, default=0)
    args, _ = parser.parse_known_args()
    run(source=args.source, output_dir=args.output_dir, decimals=args.decimals)


if __name__ == "__main__":
    main()
