#!/usr/bin/env python3
"""Create the final CMIP7 pathway SI table from the adopted production chain.

The layout matches the requested long-table style: ice-sheet section -> year ->
pathway, with median / 17–83 / 5–95 columns. 2050 is required and populated.

Created on: 2026-10-26
Author: Elizabeth H. Davis, Ph.D. (email: elizabeth.h.davis@rutgers.edu)
"""
from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd

from table_config import (
    CMIP7_CODES,
    CMIP7_OUTPUT_DIR,
    CMIP7_SOURCE,
    COMPONENTS,
    METHOD_NOTE_CMIP7,
    TABLE_YEARS,
)
from table_common import formatted_entries, load_projection_table

COMPONENT_DISPLAY = {
    "EAIS": "East Antarctic Ice Sheet (EAIS)",
    "WAIS": "West Antarctic Ice Sheet (WAIS)",
    "GrIS": "Greenland Ice Sheet (GrIS)",
    "Total AIS": "Total Antarctic Ice Sheet (AIS)",
    "Total ice sheet": "Total ice sheets",
}

HEADERS = (
    "Year and CMIP7\npathway",
    "Median\n(cm)",
    "17th–83rd\npercentile\n(cm)",
    "5th–95th\npercentile\n(cm)",
)


def build_long_table(data: pd.DataFrame, decimals: int) -> pd.DataFrame:
    lookup = data.set_index(["component", "year", "ssp"])
    rows = []
    for component in COMPONENTS:
        for year in TABLE_YEARS:
            for code in CMIP7_CODES:
                med, inner, outer = formatted_entries(
                    lookup.loc[(component, year, code)], decimals
                )
                rows.append(
                    {
                        "ice_sheet": component,
                        "year": year,
                        "cmip7_pathway": code,
                        "median_cm": med,
                        "p17_to_p83_cm": inner,
                        "p05_to_p95_cm": outer,
                    }
                )
    return pd.DataFrame(rows)


def _require_2050_populated(data: pd.DataFrame) -> None:
    rows = data.loc[data["year"].eq(2050)].copy()
    expected = len(CMIP7_CODES) * len(COMPONENTS)
    if len(rows) != expected:
        raise ValueError(
            f"CMIP7: expected {expected} complete 2050 rows, found {len(rows)}. "
            "2050 will not be left blank."
        )
    qcols = ["p05_cm", "p17_cm", "p50_cm", "p83_cm", "p95_cm"]
    if rows[qcols].isna().any().any():
        raise ValueError("CMIP7: one or more 2050 quantiles are missing.")


def make_xlsx(data: pd.DataFrame, path: Path, decimals: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "CMIP7 long table"
    ws.sheet_view.showGridLines = False

    gray = "F2F2F2"
    component_fill = "F7F7F7"
    rule = "A6A6A6"
    thin_bottom = Border(bottom=Side(style="thin", color=rule))
    medium_bottom = Border(bottom=Side(style="medium", color="000000"))

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 21
    ws.column_dimensions["D"].width = 21

    lookup = data.set_index(["component", "year", "ssp"])
    row = 1
    for comp_i, component in enumerate(COMPONENTS):
        if comp_i:
            row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row, 1, COMPONENT_DISPLAY[component])
        cell.font = Font(name="Aptos", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill("solid", fgColor=component_fill)
        ws.row_dimensions[row].height = 22
        row += 1

        for c, text in enumerate(HEADERS, start=1):
            h = ws.cell(row, c, text)
            h.font = Font(name="Aptos", size=9, bold=True)
            h.alignment = Alignment(
                horizontal="center" if c > 1 else "left",
                vertical="center",
                wrap_text=True,
            )
            h.border = medium_bottom
        ws.row_dimensions[row].height = 40
        row += 1

        for year in TABLE_YEARS:
            y = ws.cell(row, 1, str(year))
            y.font = Font(name="Aptos", size=9, bold=True)
            for c in range(1, 5):
                ws.cell(row, c).fill = PatternFill("solid", fgColor=gray)
            ws.row_dimensions[row].height = 18
            row += 1

            for code_i, code in enumerate(CMIP7_CODES):
                med, inner, outer = formatted_entries(
                    lookup.loc[(component, year, code)], decimals
                )
            
                values = (code, med, inner, outer)
            
                for c, value in enumerate(values, start=1):
                    cell = ws.cell(row, c, value)
                    cell.font = Font(name="Aptos", size=9)
                    cell.alignment = Alignment(
                        horizontal="left" if c == 1 else "center",
                        vertical="center",
                        indent=2 if c == 1 else 0,
                    )
            
                    # Only draw a rule after the LAST CMIP7 pathway
                    # in each projection year.
                    if code_i == len(CMIP7_CODES) - 1:
                        cell.border = thin_bottom
            
                ws.row_dimensions[row].height = 18
                row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    note = ws.cell(row, 1, "Note. " + METHOD_NOTE_CMIP7)
    note.font = Font(name="Aptos", size=8, italic=True)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 65

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:D{row}"
    wb.save(path)


def run(
    *,
    source: Path = CMIP7_SOURCE,
    output_dir: Path = CMIP7_OUTPUT_DIR,
    decimals: int = 0,
):
    if decimals not in range(5):
        raise ValueError("decimals must be 0, 1, 2, 3, or 4")

    data = load_projection_table(
        source, scenario_codes=CMIP7_CODES, family="CMIP7"
    )
    _require_2050_populated(data)

    output_dir = Path(output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    audit = output_dir / "CMIP7_quantiles_full_precision_cm.csv"
    data.to_csv(audit, index=False)

    long = build_long_table(data, decimals)
    csv_path = output_dir / "SI_CMIP7_long_table_all_components_final.csv"
    xlsx_path = output_dir / "SI_CMIP7_long_table_all_components_final.xlsx"
    long.to_csv(csv_path, index=False)
    make_xlsx(data, xlsx_path, decimals)

    print("\nCMIP7 final-production long SI tables complete")
    print("Source:", Path(source).expanduser().resolve())
    print("CSV:", csv_path)
    print("XLSX:", xlsx_path)
    print("2050: explicitly validated and populated")
    print("Audit table:", audit)
    return {"csv": csv_path, "xlsx": xlsx_path, "audit": audit}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=CMIP7_SOURCE)
    parser.add_argument("--output-dir", type=Path, default=CMIP7_OUTPUT_DIR)
    parser.add_argument("--decimals", type=int, default=0)
    args, _ = parser.parse_known_args()
    run(source=args.source, output_dir=args.output_dir, decimals=args.decimals)


if __name__ == "__main__":
    main()
